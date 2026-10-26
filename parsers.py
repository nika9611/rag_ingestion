"""Extract text and tables from PDF, Word, Excel, and HTML sources."""

import hashlib
import json
import re
from dataclasses import asdict
from pathlib import Path

from bs4 import BeautifulSoup
from docx import Document
from docx.table import Table
from docx.text.paragraph import Paragraph
from openpyxl import load_workbook
import pdfplumber
from schema import Block, Chunk, Kind


def sha256_file(path: Path) -> str:
    """Return the SHA-256 hex digest of the file's bytes, reading in 1 MiB parts."""
    digest = hashlib.sha256()
    with path.open("rb") as stream:
        for part in iter(lambda: stream.read(1024 * 1024), b""):
            digest.update(part)
    return digest.hexdigest()


def _sha256_record(record: Block | Chunk) -> str:
    """Hash a dataclass using a consistent UTF-8 JSON representation."""
    # Sorting keys makes metadata insertion order irrelevant. Preserve list
    # order because table rows and columns have positional meaning.
    payload = json.dumps(
        asdict(record),
        sort_keys=True,
        ensure_ascii=False,
        separators=(",", ":"),
        allow_nan=False,
    ).encode("utf-8")
    return hashlib.sha256(payload).hexdigest()


def sha256_block(block: Block) -> str:
    """Return a 64-character SHA-256 hex digest of all block fields.

    Includes document ID, kind, text, headers, rows, and metadata. Dictionary
    key order is ignored; list order and exact text (including whitespace)
    are preserved. ``Kind`` values serialize as their underlying strings.

    Metadata must contain JSON-compatible values with string dictionary keys
    and finite numbers. Unsupported values raise ``TypeError`` or
    ``ValueError``. The block is not modified and the digest is not stored.
    """
    return _sha256_record(block)


def sha256_chunk(chunk: Chunk) -> str:
    """Return a 64-character SHA-256 hex digest of all chunk fields.

    Includes chunk ID, document ID, text, and metadata. Uses the same JSON
    rules as ``sha256_block``: dictionary key order is ignored, while list
    order and exact text are preserved. Unsupported metadata raises
    ``TypeError`` or ``ValueError``.

    The digest is computed from current field values on every call. Save it
    separately to compare an object's state before and after edits.
    """
    return _sha256_record(chunk)


def document_id(path: Path, root: Path) -> str:
    """Return a document's resolved path relative to the ingestion root.

    Raises:
        ValueError: If the resolved path is outside ``root``.
    """
    return str(path.resolve().relative_to(root.resolve()))


def _clean_pdf_table(rows: list[list[str | None]]) -> list[list[str]]:
    """Remove wholly empty rows/columns while preserving meaningful cell positions."""
    matrix = [[(cell or "").strip() for cell in row] for row in rows if row is not None]
    matrix = [row for row in matrix if any(row)]
    if not matrix:
        return []
    occupied = [
        index for index in range(max(map(len, matrix)))
        if any(index < len(row) and row[index] for row in matrix)
    ]
    return [[row[index] if index < len(row) else "" for index in occupied] for row in matrix]


def parse_pdf(path: Path, doc_id: str) -> list[Block]:
    """Extract prose, numbered headings, and cleaned tables in page reading order.

    Removes repeated margin text and empty table rows/columns. The first
    nonempty row supplies headers. A headerless data-dictionary continuation
    can inherit the preceding page's headers when geometry and types agree.
    Other headerless tables still require source-specific header handling.
    Text extraction does not perform OCR on scanned pages.
    """
    output = []
    with pdfplumber.open(path) as pdf:
        # Only repeated text at the page margins is considered boilerplate.
        # Normalize page numbers so Page 1 and Page 2 match the same footer.
        def margin_key(text):
            return re.sub(r"\d+", "#", text.strip().lower())

        margins = {}
        for page_no, page in enumerate(pdf.pages, 1):
            for line in page.extract_text_lines():
                if line["top"] < page.height * 0.05 or line["top"] > page.height * 0.92:
                    margins.setdefault(margin_key(line["text"]), set()).add(page_no)
        repeated = {key for key, pages in margins.items() if len(pages) > 1}
        previous_table = None

        for page_no, page in enumerate(pdf.pages, start=1):
            found = sorted(page.find_tables(), key=lambda table: table.bbox[1])
            boxes = [table.bbox for table in found]

            def inside_table(obj):
                if obj.get("object_type") != "char":
                    return False
                x = (obj["x0"] + obj["x1"]) / 2
                y = (obj["top"] + obj["bottom"]) / 2
                return any(x0 <= x < x1 and y0 <= y <= y1 for x0, y0, x1, y1 in boxes)

            # Interleave text lines and tables instead of appending all tables
            # after the page's prose, which destroys their section context.
            events = []
            for line in page.filter(lambda obj: not inside_table(obj)).extract_text_lines():
                at_margin = line["top"] < page.height * 0.05 or line["top"] > page.height * 0.92
                if at_margin and margin_key(line["text"]) in repeated:
                    continue
                events.append((line["top"], "line", line))
            events.extend((table.bbox[1], "table", (number, table)) for number, table in enumerate(found, 1))
            events.sort(key=lambda event: event[0])
            paragraph = []
            last_bottom = None

            def flush_paragraph():
                if paragraph:
                    text = " ".join(paragraph)
                    # Group short adjacent paragraphs/list items without crossing
                    # a heading or table. Keep labels with what follows them.
                    previous = output[-1] if output else None
                    if (previous is not None and previous.kind == Kind.PARAGRAPH
                            and previous.metadata["page"] == page_no
                            and not text.endswith(":")
                            and len(previous.text) + len(text) + 2 <= 700):
                        previous.text += "\n\n" + text
                    else:
                        output.append(Block(
                            document_id=doc_id, kind=Kind.PARAGRAPH, text=text,
                            metadata={"page": page_no, "source": str(path)},
                        ))
                    paragraph.clear()

            for top, event_type, value in events:
                if event_type == "line":
                    text = value["text"].strip()
                    is_heading = bool(re.match(r"^\d+\.\s+\S", text))
                    if is_heading or (last_bottom is not None and top - last_bottom > 4):
                        flush_paragraph()
                    if is_heading:
                        output.append(Block(
                            document_id=doc_id, kind=Kind.HEADING, text=text,
                            metadata={"page": page_no, "source": str(path), "heading_level": 1},
                        ))
                    else:
                        paragraph.append(text)
                    last_bottom = value["bottom"]
                    continue

                flush_paragraph()
                last_bottom = None
                number, table = value
                rows = _clean_pdf_table(table.extract() or [])
                if not rows:
                    continue
                metadata = {"page": page_no, "table_no": number,
                            "bbox": list(table.bbox), "source": str(path)}
                headers, data = rows[0], rows[1:]
                if previous_table is not None:
                    previous, previous_height = previous_table
                    # Be conservative: inherit only for adjacent-page tables
                    # with aligned edges and a recognizable typed schema.
                    types = {"string", "integer", "float", "date", "boolean", "number", "datetime"}
                    inherited = previous.headers or []
                    same_shape = len(rows[0]) == len(inherited)
                    aligned = all(abs(table.bbox[i] - previous.metadata["bbox"][i]) < 3 for i in (0, 2))
                    typed_rows = same_shape and len(inherited) >= 2 and inherited[1].lower() in {"type", "data type"} and all(row[1].lower() in types for row in rows)
                    if (previous.metadata["page"] == page_no - 1
                            and previous.metadata["bbox"][3] > previous_height * 0.65
                            and table.bbox[1] < page.height * 0.15
                            and same_shape and aligned and typed_rows):
                        headers, data = inherited.copy(), rows
                        metadata.update({"headers_inherited": True, "continued_from_page": page_no - 1})
                block = Block(document_id=doc_id, kind=Kind.TABLE,
                              headers=headers, rows=data, metadata=metadata)
                output.append(block)
                previous_table = (block, page.height)
            flush_paragraph()
    return output


def parse_docx(path: Path, doc_id: str) -> list[Block]:
    """Extract top-level Word paragraphs, headings, and tables in source order.

    Heading levels and sequential table numbers are saved in metadata. Nested
    tables inside cells, headers, and footers are not extracted by this parser.
    The first table row is assumed to be the column header.
    """
    result = []
    doc = Document(path)
    table_no = 0

    for element in doc.iter_inner_content():

        # Process paragraphs and headings
        if isinstance(element, Paragraph):
            text = element.text.strip()

            if text:
                style_name = element.style.name if element.style else "Normal"

                kind = (
                    Kind.HEADING
                    if style_name.lower().startswith("heading")
                    else Kind.PARAGRAPH
                )

                metadata = {
                    "style": style_name,
                    "source": str(path),
                }
                if kind == Kind.HEADING:
                    match = re.match(r"heading\s+(\d+)", style_name, re.I)
                    metadata["heading_level"] = int(match.group(1)) if match else 1

                result.append(
                    Block(
                        document_id=doc_id,
                        kind=kind,
                        text=text,
                        metadata=metadata,
                    )
                )

        # Process tables
        elif isinstance(element, Table):
            table_no += 1
            rows = [
                [cell.text.strip() for cell in row.cells]
                for row in element.rows
            ]

            if rows:
                result.append(
                    Block(
                        document_id=doc_id,
                        kind=Kind.TABLE,
                        headers=rows[0],
                        rows=rows[1:],
                        metadata={
                            "table_no": table_no,
                            "source": str(path),
                        },
                    )
                )

    return result


def parse_xlsx(
    path: Path,
    doc_id: str,
    data_only: bool = True
) -> list[Block]:
    """Extract contiguous, nonempty worksheet regions as table blocks.

    Blank rows separate regions, and column offsets and merged ranges are
    retained in metadata. ``data_only=True`` reads cached formula results,
    which can be missing when a workbook has not been recalculated in Excel.
    The first row of a region is assumed to contain headers unless it is a
    merged title row spanning the region's columns.

    This heuristic may split a table containing intentional blank rows.
    Normal worksheet loading preserves merged-range metadata but uses memory
    proportional to the size of the workbook.
    """

    blocks = []

    wb = load_workbook(
        path,
        data_only=data_only,
        read_only=False
    )

    try:
        for ws in wb.worksheets:

            sections = []
            current_section = []

            # Identify contiguous nonempty row groups
            for row_no, row in enumerate(
                ws.iter_rows(),
                start=1
            ):
                values = [
                    "" if cell.value is None
                    else str(cell.value).strip()
                    for cell in row
                ]

                if any(values):
                    current_section.append((row_no, values))

                elif current_section:
                    sections.append(current_section)
                    current_section = []

            if current_section:
                sections.append(current_section)

            # Convert each region into a table block
            for table_no, section in enumerate(
                sections,
                start=1
            ):

                # Determine occupied column boundaries
                occupied_cols = [
                    col_no
                    for _, row in section
                    for col_no, value in enumerate(row)
                    if value
                ]

                first_col = min(occupied_cols)
                last_col = max(occupied_cols) + 1

                matrix = [
                    row[first_col:last_col]
                    for _, row in section
                ]

                # A common report pattern uses a merged title above headers.
                table_title = None
                header_row = section[0][0]
                if len(matrix) >= 2:
                    first_row = section[0][0]
                    title_is_merged = any(
                        merged.min_row == merged.max_row == first_row
                        and merged.min_col <= first_col + 1
                        and merged.max_col >= last_col
                        for merged in ws.merged_cells.ranges
                    )
                    if (
                        title_is_merged
                        and sum(bool(value) for value in matrix[0]) == 1
                        and sum(bool(value) for value in matrix[1]) >= 2
                    ):
                        table_title = next(value for value in matrix[0] if value)
                        matrix = matrix[1:]
                        header_row += 1

                blocks.append(
                    Block(
                        document_id=doc_id,
                        kind=Kind.TABLE,
                        headers=matrix[0],
                        rows=matrix[1:],
                        metadata={
                            "sheet": ws.title,
                            "table_no": table_no,
                            "start_row": section[0][0],
                            "header_row": header_row,
                            "end_row": section[-1][0],
                            "start_col": first_col + 1,
                            "end_col": last_col,
                            "table_title": table_title,
                            "merged_ranges": [
                                str(r)
                                for r in ws.merged_cells.ranges
                            ],
                            "source": str(path),
                            "formula_mode": (
                                "cached_values"
                                if data_only else "formulas"
                            )
                        }
                    )
                )

    finally:
        wb.close()

    return blocks


def parse_html(path: Path, doc_id: str) -> list[Block]:
    """Extract visible HTML text after removing common boilerplate tags.

    Tables are flattened into prose; use a dedicated HTML table parser when
    retaining row and column structure is important.
    """
    soup = BeautifulSoup(path.read_text(encoding="utf-8"), "html.parser")

    for tag in soup.select("script, style, nav, footer"):
        tag.decompose()
    text = soup.get_text(" ", strip=True)

    if not text:
        return []

    return [
        Block(
            document_id=doc_id,
            kind=Kind.PARAGRAPH,
            text=text,
            metadata={"source": str(path)},
        )
    ]


def ocr_pdf_page(path: Path, page_no: int) -> str:
    """OCR one 1-based PDF page at 250 DPI using Poppler and Tesseract.

    Requires the optional ``pdf2image`` and ``pytesseract`` packages plus their
    system executables. This helper does not run automatically in parse_pdf.
    """
    from pdf2image import convert_from_path
    import pytesseract

    images = convert_from_path(
        str(path),
        dpi=250,
        first_page=page_no,
        last_page=page_no
    )
    return pytesseract.image_to_string(images[0]) if images else ""
